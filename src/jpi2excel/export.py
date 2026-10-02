# Copyright © 2025-2026 by Alan M. Marcum, Nescorna Professional.
# All rights reserved.

"""Revision: 9. Shared CSV/XLSX layout, metadata, units and alarm reporting."""
from copy import copy
import csv
from datetime import datetime, timedelta
import errno
from pathlib import Path
import re

from .model import ConversionError

NAVIGATION = ('LAT', 'LNG', 'ALT', 'SPD')
LABELS = {'OILP': 'Oil P', 'OILT': 'Oil T', 'BAT': 'Batt', 'USD': 'Used',
          'DIF': 'Diff', 'CLD': 'Cold', 'HP': '% HP', 'MARK': 'Mark',
          'LAT': 'Lat', 'LNG': 'Lon', 'ALT': 'Alt', 'SPD': 'Speed'}

# Opening geometry from data/t2.xlsx, with width from data/t4.xlsx (twips).
# Encode the reference values so exporting does not depend on that sample file.
WINDOW_GEOMETRY = {'xWindow': 4280, 'yWindow': 2700,
                   'windowWidth': 36860, 'windowHeight': 18380}
NARROW_SENSORS = {'FF', 'HP', 'OAT', 'CDT', 'IAT', 'OILP', 'OILT', 'BAT',
                  'USD', 'MARK', 'DIF', 'CLD'}

# XLSX stores widths in character units, not pixels. Calibrate to the supplied
# macOS Calibri 11 workbook: width 20 ~ 120 px, width 75 ~ 450 px.
# Quantize to Excel's 1/256-character precision; rendering varies by font/viewer.
def excel_width(pixels):
    return int(pixels * 256 / 6) / 256


def flight_width(code):
    if code in NARROW_SENSORS or re.fullmatch(r'C\d+', code):
        return 53
    return {'Sample': 50, 'DeltaT': 44, 'DateTime': 115, 'Limits': 125}.get(code, 55)


def ordered_codes(flight):
    present = set(flight.series)
    def group(prefix):
        return sorted((code for code in present if re.fullmatch(prefix + r'\d+', code)),
                      key=lambda code: int(code[1:]))
    order = ['MAP', 'RPM', *group('E'), *group('T'), *group('C'),
             'FF', 'HP', 'OAT', 'CDT', 'IAT', 'OILP', 'OILT', 'BAT',
             'USD', 'MARK', 'DIF', 'CLD']
    ordered = [code for code in order if code in present]
    ordered += sorted(present - set(ordered) - set(NAVIGATION))
    return ordered + ['Limits'] + [code for code in NAVIGATION if code in present]


def label(code, flight):
    if re.fullmatch(r'[ECT]\d+', code):
        category = {'E': 'EGT', 'C': 'CHT', 'T': 'TIT'}[code[0]]
        return 'TIT' if code[0] == 'T' and flight.tit_count == 1 else f'{category} {code[1:]}'
    return LABELS.get(code, code)


def units(code, download):
    engine = download.metadata['Engine temperature units']
    if re.fullmatch(r'[ECT]\d+', code) or code in ('IAT', 'CDT', 'OILT', 'DIF'):
        return engine
    return {'OAT': download.metadata['OAT temperature units'], 'CLD': f'{engine}/min',
            'FF': download.metadata['Fuel flow units'], 'USD': download.metadata['Fuel units'],
            'MAP': 'inHg', 'RPM': 'rpm', 'BAT': 'V', 'OILP': 'psi', 'HP': '%',
            'LAT': 'degrees', 'LNG': 'degrees', 'ALT': 'ft', 'SPD': 'knots'}.get(code, '')


def alarm_category(code):
    return 'CHT' if re.fullmatch(r'C\d+', code) else 'TIT' if re.fullmatch(r'T\d+', code) else code


def alarm_properties(download):
    for code, (low, high) in download.alarms.items():
        unit = 'V' if code == 'BAT' else download.metadata['Engine temperature units']
        if code == 'CLD':
            unit += '/min'
        for direction, value in (('low', low), ('high', high)):
            if value is not None:
                yield code, direction, f'{code} {direction} alarm ({unit})', value


def limits(flight, download, index):
    violations = []
    for code in ordered_codes(flight):
        if code == 'Limits':
            continue
        value = flight.series[code][index]
        if value is None:
            continue
        category = alarm_category(code)
        low, high = download.alarms.get(category, (None, None))
        # Use individual cylinder/TIT labels and JPI codes for other alarms.
        name = label(code, flight) if category in ('CHT', 'TIT') else code
        if low is not None and value < low:
            violations.append(name + ('-L' if high is not None else ''))
        if high is not None and value > high:
            violations.append(name + ('-H' if low is not None else ''))
    return ';'.join(violations)


def headers(flight):
    return ['Sample', 'Delta T', 'DateTime'] + [label(code, flight) for code in ordered_codes(flight)]


def rows(flight, download):
    codes = ordered_codes(flight)
    for index, elapsed in enumerate(flight.elapsed):
        yield [index + 1, elapsed, flight.start + timedelta(seconds=elapsed)] + [
            limits(flight, download, index) if code == 'Limits' else flight.series[code][index]
            for code in codes]


def duration_text(seconds):
    hours, remaining = divmod(int(seconds), 3600)
    minutes, seconds = divmod(remaining, 60)
    return f'{hours:02}:{minutes:02}:{seconds:02}'


def info_items(download, minimum_minutes, *, flight_ids=None):
    yield from download.metadata.items()
    yield 'Minimum duration (minutes)', minimum_minutes
    yield 'Flights meeting minimum duration', sum(f.duration >= minimum_minutes * 60 for f in download.flights)
    for _, _, name, value in alarm_properties(download):
        yield name, value
    for flight in download.flights:
        if flight_ids is not None and flight.id not in flight_ids:
            continue
        prefix = f'Flight {flight.id} '
        properties = {'start': flight.start, 'duration': duration_text(flight.duration),
                      'samples': flight.samples, 'recording interval (seconds)': flight.interval,
                      'rated HP': flight.horsepower, 'engine': flight.engine,
                      'cylinder count': flight.cylinder_count, 'TIT channel count': flight.tit_count,
                      'IAT present': 'Yes' if 'IAT' in flight.series else 'No',
                      'CDT present': 'Yes' if 'CDT' in flight.series else 'No',
                      'navigation': 'GPS' if flight.gps else 'No GPS',
                      'fuel unit code': flight.fuel_unit,
                      'configuration words': ', '.join(map(str, flight.configuration)),
                      'channels': ', '.join(flight.series),
                      'column units': '; '.join(f'{label(c, flight)}={units(c, download)}'
                                               for c in flight.series if units(c, download))}
        for name, value in properties.items():
            yield prefix + name, value
    for index, header in enumerate(download.raw_headers, 1):
        yield f'Raw header {index}', header
    for index, warning in enumerate(download.warnings, 1):
        yield f'Warning {index}', warning


def write_csv(path, flight, download):
    with Path(path).open('x', encoding='utf-8', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(headers(flight))
        for row in rows(flight, download):
            row[2] = row[2].isoformat(sep=' ')
            writer.writerow(row)


def format_alarm_columns(sheet, flight, download, references):
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not flight.samples:
        return
    colors = {'high': ('FFFFC7CE', 'FF9C0006', 'greaterThanOrEqual'),
              'low': ('FFFFEB9C', 'FF9C6500', 'lessThanOrEqual')}
    for index, code in enumerate(ordered_codes(flight), 4):
        if code == 'Limits':
            continue
        available = [(direction, references[(alarm_category(code), direction)])
                     for direction in colors if (alarm_category(code), direction) in references]
        if not available:
            continue
        column = get_column_letter(index)
        cells = f'{column}2:{column}{flight.samples + 1}'
        # A blank reading must not be compared as zero against a low threshold.
        sheet.conditional_formatting.add(cells, FormulaRule(
            formula=[f'NOT(ISNUMBER({column}2))'], stopIfTrue=True))
        for direction, reference in available:
            fill, text, operator = colors[direction]
            sheet.conditional_formatting.add(cells, CellIsRule(
                operator=operator, formula=[reference],
                # Differential fills use the background color, as in Excel's
                # standard red/yellow presets; do not add a foreground pattern.
                fill=PatternFill(bgColor=fill), font=Font(color=text)))
        low, high = download.alarms[alarm_category(code)]
        if any(value is not None and ((high is not None and value >= high)
                                     or (low is not None and value <= low))
               for value in flight.series[code]):
            header = sheet.cell(1, index)
            font = copy(header.font)
            font.color = 'FFFF0000'
            header.font = font
            sheet.sheet_properties.tabColor = 'FFFFFF00'


def write_xlsx(path, selected, downloads, minimum_minutes, summary=True, graphs=False):
    """Write flight data; summary=False hides the Summary used by alarm rules."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for flight, _ in selected:
        if flight.samples > 1_048_575:
            raise ConversionError(f'Flight {flight.id} exceeds Excel row limit; use CSV', errno.EFBIG)
    workbook = Workbook()
    for name, value in WINDOW_GEOMETRY.items():
        setattr(workbook.views[0], name, value)
    workbook.remove(workbook.active)

    def style(sheet, pixel_widths):
        sheet.freeze_panes = 'C2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='305070')
        for column, pixels in enumerate(pixel_widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = excel_width(pixels)

    row_counts = {}

    def append(sheet, values):
        sheet.append(values)
        row_counts[sheet.title] = row_counts.get(sheet.title, 0) + 1
        row_number = row_counts[sheet.title]
        # Metadata are source text, even when an aircraft ID begins with '='.
        for column in range(1, len(values) + 1):
            cell = sheet.cell(row_number, column)
            if isinstance(cell.value, str):
                cell.data_type = 's'
            elif isinstance(cell.value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss'

    summary_sheet = workbook.create_sheet('Summary')

    used = set()
    graph_warnings = []
    flight_sheets = []
    for flight, download in selected:
        title = f'Flight {flight.id}'
        suffix = 1
        while title.casefold() in used:
            suffix += 1
            title = f'Flight {flight.id} ({suffix})'
        used.add(title.casefold())
        sheet = workbook.create_sheet(title)
        append(sheet, headers(flight))
        for row in rows(flight, download):
            append(sheet, row)
        codes = ['Sample', 'DeltaT', 'DateTime', *ordered_codes(flight)]
        style(sheet, [flight_width(code) for code in codes])
        flight_sheets.append((sheet, flight, download))
        if graphs:
            from .graphs import add_graphs
            warnings = add_graphs(workbook, sheet, flight, download,
                                  {code: column for column, code in enumerate(codes, 1)})
            graph_warnings.extend((download, warning) for warning in warnings)
    sheet = summary_sheet
    alarm_references = {}
    append(sheet, ['Source', 'Property', 'Value'])
    for download in downloads:
        properties = {name: (code, direction) for code, direction, name, _ in alarm_properties(download)}
        references = alarm_references.setdefault(id(download), {})
        for name, value in info_items(download, minimum_minutes):
            append(sheet, [download.path.name, name, value])
            if name in properties:
                references[properties[name]] = f'Summary!$C${row_counts[sheet.title]}'
        count = sum(d is download for _, d in selected)
        append(sheet, [download.path.name, 'Exported flights', count])
        for owner, warning in graph_warnings:
            if owner is download:
                append(sheet, [download.path.name, 'Graph warning', warning])
    style(sheet, [85, 225, 450])
    wrapped = Alignment(wrap_text=True)
    for cell in sheet['C']:
        cell.alignment = wrapped
    for flight_sheet, flight, download in flight_sheets:
        format_alarm_columns(flight_sheet, flight, download, alarm_references[id(download)])
    if not summary:
        summary_sheet.sheet_state = 'hidden'
        workbook.active = flight_sheets[0][0]
    with Path(path).open('xb') as stream:
        workbook.save(stream)
    workbook.close()
    return [f'{download.path.name}: {warning}' for download, warning in graph_warnings]
