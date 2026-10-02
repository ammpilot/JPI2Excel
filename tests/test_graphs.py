# Copyright © 2025-2026 by Alan M. Marcum, Nescorna Professional.
# All rights reserved.

"""Revision: 4. Graph structure, numeric time, bounds, styling and CLI regression tests."""
from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
import errno
import io
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from jpi2excel.cli import main
from jpi2excel.export import write_xlsx
from jpi2excel.reader import read_jpi
from test_parser import synthetic_file

NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


class GraphTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        self.input = self.directory / 'source.JPI'
        self.input.write_bytes(synthetic_file())
        self.download = read_jpi(self.input)
        self.flight = self.download.flights[0]
        self.flight.elapsed = [0, 2, 3, 5, 7]
        self.flight.series = {
            **{f'E{i}': [700, 1400, None, 1551, 1500] for i in range(1, 9)},
            **{f'C{i}': [201, 300, None, 401, 300] for i in range(1, 9)},
            'T1': [800, 1450, 1610, 1751, 1500],
            'OILT': [95, 180, 200, 451, 150],
            'FF': [0, 12, 15.2, 10, 0], 'RPM': [699, 2400, 2611, 2500, 800],
        }

    def invoke(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = main(list(map(str, args)))
        return code, out.getvalue(), err.getvalue()

    def export(self, **kwargs):
        target = self.directory / 'graphs.xlsx'
        warnings = write_xlsx(target, [(self.flight, self.download)], [self.download],
                              0, graphs=True, **kwargs)
        book = load_workbook(target)
        self.addCleanup(book.close)
        return book, warnings

    def test_axes_series_styles_and_numeric_time(self):
        book, warnings = self.export()
        self.assertEqual(warnings, [])
        self.assertEqual(book.sheetnames, ['Summary', 'Flight 415', 'Graph 415'])
        exhaust, cylinder = book['Graph 415']._charts
        for chart, low, high, secondary_low, secondary_high in (
                (exhaust, 700, 1900, 1, 17), (cylinder, 0, 500, 600, 2800)):
            self.assertEqual(chart.scatterStyle, 'line')
            self.assertEqual((chart.x_axis.scaling.min, chart.x_axis.scaling.max), (0, 7))
            self.assertEqual((chart.y_axis.scaling.min, chart.y_axis.scaling.max), (low, high))
            secondary = chart._charts[1]
            self.assertEqual((secondary.y_axis.scaling.min, secondary.y_axis.scaling.max),
                             (secondary_low, secondary_high))
            self.assertEqual(secondary.y_axis.axPos, 'r')
            self.assertEqual(secondary.x_axis.scaling.max, 7)
            self.assertTrue(secondary.x_axis.delete)
            self.assertIsNone(secondary.x_axis.title)
            self.assertIsNone(secondary.x_axis.tickLblPos)  # openpyxl's representation of "none"
            self.assertEqual(chart.x_axis.numFmt.formatCode, '#,##0')
            self.assertFalse(chart.x_axis.numFmt.sourceLinked)
            self.assertEqual(secondary.series[0].graphicalProperties.line.prstDash, 'sysDash')
            self.assertEqual(chart.legend.position, 'b')
            for series in chart.series[:-1] + secondary.series:
                self.assertEqual(series.xVal.numRef.f, "'Flight 415'!$B$2:$B$6")
                self.assertEqual([p.v for p in series.xVal.numRef.numCache.pt], [0, 2, 3, 5, 7])
                self.assertEqual(series.marker.symbol, None)
                self.assertFalse(series.smooth)
            limit = chart.series[-1]
            self.assertEqual(limit.graphicalProperties.line.solidFill.srgbClr, 'FF0000')
            self.assertEqual(limit.graphicalProperties.line.w, 19050)  # 2 px at 96 dpi
            self.assertEqual(limit.graphicalProperties.line.prstDash, 'solid')
            self.assertEqual([p.v for p in limit.xVal.numLit.pt], [0, 7])
        for chart, value, name in [(exhaust, 1700, 'TIT Limit'), (cylinder, 400, 'CHT Limit')]:
            self.assertEqual(chart.series[-1].tx.v, name)
            self.assertEqual([p.v for p in chart.series[-1].yVal.numLit.pt], [value, value])
        colors = [s.graphicalProperties.line.solidFill.srgbClr for s in exhaust.series[:8]]
        self.assertEqual(colors[:6], ['FFC000', '00B050', '00B0F0', '0070C0', '7030A0', '945200'])
        self.assertEqual(len(set(colors)), 8)
        self.assertEqual(colors, [s.graphicalProperties.line.solidFill.srgbClr for s in cylinder.series[:8]])
        # Gap retains its sample index in the cache, rather than becoming zero.
        self.assertEqual([p.idx for p in exhaust.series[0].yVal.numRef.numCache.pt], [0, 1, 3, 4])
        with ZipFile(self.directory / 'graphs.xlsx') as archive:
            xml = ET.fromstring(archive.read('xl/charts/chart1.xml'))
            fonts = {n.attrib['typeface'] for n in xml.findall('.//a:latin', NS)}
            self.assertTrue({'Cambria', 'Calibri'} <= fonts)
            self.assertEqual(len(xml.findall('.//c:valAx', NS)), 4)
            self.assertFalse(xml.findall('.//c:catAx', NS))
            for name in ('xl/charts/chart1.xml', 'xl/charts/chart2.xml'):
                xml = ET.fromstring(archive.read(name))
                time_titles = [node.text for node in xml.findall('.//c:valAx/c:title//a:t', NS)
                               if node.text == 'Time (seconds)']
                self.assertEqual(time_titles, ['Time (seconds)'])

    def test_omitted_series_and_thresholds(self):
        self.flight.series = {'E3': [1200] * 5, 'C3': [300] * 5}
        self.download.alarms = {}
        book, warnings = self.export(summary=False)
        self.assertEqual(warnings, [])
        self.assertEqual(book.sheetnames, ['Summary', 'Flight 415', 'Graph 415'])
        self.assertEqual(book['Summary'].sheet_state, 'hidden')
        for chart in book['Graph 415']._charts:
            self.assertEqual(len(chart._charts), 1)
            self.assertEqual(len(chart.series), 1)
            self.assertEqual(chart.series[0].graphicalProperties.line.solidFill.srgbClr, '00B0F0')
        exhaust, cylinder = book['Graph 415']._charts
        self.assertEqual((exhaust.y_axis.scaling.min, exhaust.y_axis.scaling.max), (1200, 1300))
        self.assertEqual((cylinder.y_axis.scaling.min, cylinder.y_axis.scaling.max), (300, 350))

    def test_alarms_control_maximum_and_negative_minimum_rounds_down(self):
        self.flight.series = {'E1': [-1, 1200, None, 1300, 1200], 'C1': [-1, 300, None, 300, 300]}
        self.download.metadata['Engine temperature units'] = 'C'
        book, _ = self.export()
        exhaust, cylinder = book['Graph 415']._charts
        self.assertEqual((exhaust.y_axis.scaling.min, exhaust.y_axis.scaling.max), (-100, 1800))
        self.assertEqual((cylinder.y_axis.scaling.min, cylinder.y_axis.scaling.max), (-100, 450))
        for chart in (exhaust, cylinder):
            self.assertEqual(chart.y_axis.title.tx.rich.p[0].r[0].t, 'Temperature, ºC')

    def test_missing_temperatures_omit_chart_and_warn(self):
        self.flight.series = {'E1': [None] * 5, 'RPM': [2400] * 5, 'C1': [300] * 5}
        book, warnings = self.export()
        self.assertEqual(len(book['Graph 415']._charts), 1)
        self.assertEqual(len(warnings), 1)
        self.assertIn('Exhaust Temperatures', warnings[0])
        self.assertIn('no temperature data', warnings[0])
        self.assertTrue(any('no temperature data' in str(r[2].value) for r in book['Summary']))

    def test_zero_duration_and_zero_flow_have_valid_axes(self):
        self.flight.elapsed = [0]
        self.flight.series = {'E1': [1000], 'C1': [100], 'FF': [0]}
        book, _ = self.export()
        chart = book['Graph 415']._charts[0]
        self.assertGreater(chart.x_axis.scaling.max, chart.x_axis.scaling.min)
        axis = chart._charts[1].y_axis
        self.assertGreater(axis.scaling.max, axis.scaling.min)

    def test_duplicate_flight_names_pair_with_graphs(self):
        other = deepcopy(self.download)
        target = self.directory / 'duplicates.xlsx'
        write_xlsx(target, [(self.flight,self.download),(other.flights[0],other)],
                   [self.download,other], 0, graphs=True)
        book = load_workbook(target)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames,
                         ['Summary','Flight 415','Graph 415','Flight 415 (2)','Graph 415 (2)'])
        chart = book['Graph 415 (2)']._charts[0]
        self.assertEqual(chart.series[0].xVal.numRef.f, "'Flight 415 (2)'!$B$2:$B$6")

    def test_cli_aliases_actions_and_selection(self):
        fixture = Path(__file__).resolve().parents[1] / 'testdata/U260919.JPI'
        for flag, action in [('--graph','--xls'), ('--graphs','--xls-separate')]:
            target = self.directory / (action[2:] + '.xlsx')
            code, out, err = self.invoke('--info', action, flag, '--flights','415',
                                         '--output',target,fixture)
            self.assertEqual(code, 0, err)
            self.assertIn('Flight 415 samples: 724', out)
            book = load_workbook(target)
            self.assertEqual(book.sheetnames[-2:], ['Flight 415','Graph 415'])
            self.assertNotIn('Flight 416', book.sheetnames)
            self.assertEqual(len(book['Graph 415']._charts), 2)
            book.close()
        for flag in ('--graph','--graphs'):
            code, _, err = self.invoke('--info','--csv',flag,self.input)
            self.assertEqual(code, errno.EINVAL)
            self.assertIn('--csv', err)
            for action in ('--info','--list','--validate'):
                self.assertEqual(self.invoke(action,flag,self.input)[0], 0)
        self.assertIn('--graphs', self.invoke('--help')[1])

    def test_cli_reports_chart_omissions_without_failing(self):
        target = self.directory / 'missing.xlsx'
        code, _, err = self.invoke('--xls-separate','--graphs','--minimum-duration','0',
                                    '--output',target,self.input)
        self.assertEqual(code, 0, err)
        self.assertIn('omitted Cylinder Temperatures graph: no temperature data', err)
