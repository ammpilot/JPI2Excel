# Copyright © 2025-2026 by Alan M. Marcum, Nescorna Professional.
# All rights reserved.

"""Revision: 4. Excel alarm rules, source-specific references, and missing values."""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from jpi2excel.export import write_xlsx
from jpi2excel.reader import read_jpi
from test_parser import synthetic_file


class ConditionalFormattingTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        source = self.directory / 'source.JPI'
        source.write_bytes(synthetic_file())
        self.download = read_jpi(source)
        self.flight = self.download.flights[0]

    def export(self, selected=None, downloads=None, **kwargs):
        path = self.directory / 'conditional.xlsx'
        write_xlsx(path, selected or [(self.flight, self.download)],
                   downloads or [self.download], 0, **kwargs)
        book = load_workbook(path)
        self.addCleanup(book.close)
        return book, path

    @staticmethod
    def rules(sheet):
        return {str(item.sqref): sheet.conditional_formatting[item]
                for item in sheet.conditional_formatting}

    @staticmethod
    def alarm_refs(summary, source):
        return {row[1].value: f'Summary!$C${row[2].row}' for row in summary
                if row[0].value == source}

    def test_every_available_limit_has_classic_rule_and_standard_colors(self):
        book, path = self.export()
        sheet = book['Flight 415']
        refs = self.alarm_refs(book['Summary'], self.download.path.name)
        rules = self.rules(sheet)
        expected = {'Batt': ('BAT', 'V', ('low','high')),
                    'Oil T': ('OILT', 'F', ('low','high')),
                    'TIT': ('TIT', 'F', ('high',)),
                    'Diff': ('DIF', 'F', ('high',)),
                    'Cold': ('CLD', 'F/min', ('low',)),
                    **{f'CHT {n}': ('CHT', 'F', ('high',)) for n in range(1, 7)}}
        self.assertEqual(len(rules), len(expected))
        for cell in sheet[1]:
            if cell.value not in expected:
                self.assertNotIn(f'{cell.column_letter}2:{cell.column_letter}6', rules)
                continue
            category, unit, directions = expected[cell.value]
            column_rules = rules[f'{cell.column_letter}2:{cell.column_letter}6']
            # Excel treats blanks as zero in numeric comparisons; skip them first.
            blank_rule = column_rules[0]
            self.assertEqual(blank_rule.type, 'expression')
            self.assertEqual(blank_rule.formula, [f'NOT(ISNUMBER({cell.column_letter}2))'])
            self.assertTrue(blank_rule.stopIfTrue)
            self.assertIsNone(blank_rule.dxf)
            comparisons = {rule.operator: rule for rule in column_rules[1:]}
            self.assertEqual(len(comparisons), len(directions))
            for direction in directions:
                operator = 'lessThanOrEqual' if direction == 'low' else 'greaterThanOrEqual'
                rule = comparisons[operator]
                self.assertEqual(rule.type, 'cellIs')
                self.assertEqual(rule.formula, [refs[f'{category} {direction} alarm ({unit})']])
                fill, font = ('FFFFEB9C','FF9C6500') if direction == 'low' else ('FFFFC7CE','FF9C0006')
                self.assertEqual(rule.dxf.fill.bgColor.rgb, fill)
                self.assertIsNone(rule.dxf.fill.patternType)
                self.assertEqual(rule.dxf.font.color.rgb, font)
        # The saved file contains cell comparisons with references, not constants.
        with ZipFile(path) as archive:
            root = ET.fromstring(archive.read('xl/worksheets/sheet2.xml'))
            ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            for node in root.findall('.//s:cfRule[@type="cellIs"]/s:formula', ns):
                self.assertTrue(node.text.startswith('Summary!$C$'))
            # Preset backgrounds must survive serialization into differential
            # styles; a foreground pattern alone leaves Excel's background empty.
            styles = ET.fromstring(archive.read('xl/styles.xml'))
            backgrounds = []
            for pattern in styles.findall('s:dxfs/s:dxf/s:fill/s:patternFill', ns):
                self.assertIsNone(pattern.find('s:fgColor', ns))
                backgrounds.append(pattern.find('s:bgColor', ns).get('rgb'))
            self.assertEqual(set(backgrounds), {'FFFFC7CE', 'FFFFEB9C'})

    def test_multi_input_rules_reference_each_sources_own_alarms(self):
        other = deepcopy(self.download)
        other.path = self.directory / 'other.JPI'
        other.alarms['OILT'] = (100, 230)
        book, _ = self.export([(self.flight,self.download),(other.flights[0],other)],
                              [self.download,other])
        references = []
        for title, download in [('Flight 415',self.download),('Flight 415 (2)',other)]:
            sheet = book[title]
            col = next(cell.column_letter for cell in sheet[1] if cell.value == 'Oil T')
            rules = self.rules(sheet)[f'{col}2:{col}6']
            high = next(rule for rule in rules if rule.operator == 'greaterThanOrEqual')
            expected = self.alarm_refs(book['Summary'], download.path.name)['OILT high alarm (F)']
            self.assertEqual(high.formula, [expected])
            references.append(high.formula[0])
        self.assertNotEqual(*references)

    def test_omitted_sensors_and_missing_thresholds_do_not_create_rules(self):
        self.flight.series = {'E1': [250] * 5, 'OILT': [None] * 5}
        self.download.alarms = {'CHT': (None, 400), 'OILT': (None, None)}
        book, _ = self.export()
        self.assertEqual(len(book['Flight 415'].conditional_formatting), 0)

    def test_separate_workbook_uses_hidden_summary_and_opens_on_flight(self):
        book, _ = self.export(summary=False, graphs=True)
        self.assertEqual(book.sheetnames, ['Summary', 'Flight 415', 'Graph 415'])
        self.assertEqual(book['Summary'].sheet_state, 'hidden')
        self.assertEqual(book.active.title, 'Flight 415')
        refs = self.alarm_refs(book['Summary'], self.download.path.name)
        sheet = book['Flight 415']
        column = next(cell.column_letter for cell in sheet[1] if cell.value == 'Oil T')
        rules = self.rules(sheet)[f'{column}2:{column}6']
        high = next(rule for rule in rules if rule.operator == 'greaterThanOrEqual')
        self.assertEqual(high.formula, [refs['OILT high alarm (F)']])

    def test_tab_and_headers_flag_inclusive_limits_for_individual_channels(self):
        self.flight.series = {'C1': [399,400,None,300,300], 'C2': [399] * 5,
                              'T1': [1700] * 5, 'T2': [1699] * 5,
                              'BAT': [24,26,27,28,None], 'OILT': [96] * 5,
                              'CLD': [-199,-200,None,-100,0], 'MAP': [30] * 5}
        book, _ = self.export()
        sheet = book['Flight 415']
        self.assertEqual(sheet.sheet_properties.tabColor.rgb, 'FFFFFF00')
        flagged = {'CHT 1', 'TIT 1', 'Batt', 'Cold'}
        for header in sheet[1]:
            self.assertEqual(header.font.color.rgb,
                             'FFFF0000' if header.value in flagged else '00FFFFFF')
            self.assertTrue(header.font.bold)
            self.assertEqual(header.fill.fgColor.rgb, '00305070')
        self.assertIsNone(book['Summary'].sheet_properties.tabColor)

    def test_missing_and_in_range_values_leave_tab_and_headers_unmarked(self):
        self.flight.series = {'C1': [399,None,399,399,399], 'T1': [None] * 5,
                              'OILT': [96,100,120,180,219], 'BAT': [24.1] * 5,
                              'CLD': [-199] * 5, 'DIF': [499] * 5}
        book, _ = self.export(summary=False)
        sheet = book['Flight 415']
        self.assertIsNone(sheet.sheet_properties.tabColor)
        self.assertTrue(all(header.font.color.rgb == '00FFFFFF' for header in sheet[1]))
