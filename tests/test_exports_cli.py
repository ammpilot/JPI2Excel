"""Revision: 8. Export contracts and complete command-line behavior."""
from contextlib import redirect_stdout, redirect_stderr
import csv
import errno
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from jpi2excel.cli import main
from jpi2excel.export import headers, limits, rows, write_csv, write_xlsx
from jpi2excel.reader import read_jpi
from test_parser import synthetic_file, dollar, FIXTURES


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.input = self.directory / 'a.JPI'
        self.input.write_bytes(synthetic_file())
        self.download = read_jpi(self.input)
        self.flight = self.download.flights[0]

    def invoke(self, *args):
        stdout, stderr = io.StringIO(), io.StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            code = main(list(map(str, args)))
        return code, stdout.getvalue(), stderr.getvalue()

    def test_navigation_last_and_labels(self):
        for code in ('SPD', 'LAT', 'LNG', 'ALT'):
            self.flight.series[code] = [None] * 5
        cols = headers(self.flight)
        self.assertEqual(cols[:7], ['Sample','Delta T','DateTime','MAP','RPM','EGT 1','EGT 2'])
        self.assertEqual(cols[-5:], ['Limits', 'Lat', 'Lon', 'Alt', 'Speed'])
        self.assertIn('TIT', cols)
        self.assertNotIn('TIT 1', cols)
        self.assertIn('Batt', cols)
        self.assertIn('Cold', cols)
        self.assertIn('Oil P', cols)
        self.assertIn('Oil T', cols)
        self.assertIn('% HP', cols)
        self.assertIn('Mark', cols)

    def test_verbose_counts_for_original_four_flight_input(self):
        target = self.directory / 'verbose.xlsx'
        code, _, err = self.invoke('--verbose','--xls','--output',target,FIXTURES/'U260919.JPI')
        self.assertEqual(code, 0, err)
        self.assertIn('Exported 3 flights to 1 file\n', err)
        self.assertIn('Skipped 1 flight below minimum duration\n', err)
        self.assertIn('U260919.JPI: 4 flights\n', err)
        self.assertIn('Flight 415: 6 cylinders; 1 TIT channel; duration 00:24:06\n', err)

    def test_verbose_singular_configuration_and_zero_channel_count(self):
        self.flight.series = {'E1': [250] * 5, 'C1': [300] * 5}
        target = self.directory / 'single.xlsx'
        with patch('jpi2excel.reader.read_jpi', return_value=self.download):
            code, _, err = self.invoke('--verbose','--xls','--minimum-duration','0',
                                       '--output',target,self.input)
        self.assertEqual(code, 0, err)
        self.assertIn('Exported 1 flight to 1 file\n', err)
        self.assertIn('Skipped 0 flights below minimum duration\n', err)
        self.assertIn('a.JPI: 1 flight\n', err)
        self.assertIn('Flight 415: 1 cylinder; 0 TIT channels;', err)

    def test_absent_channel_omitted_declared_blank_retained(self):
        self.flight.series.pop('CDT')
        self.assertNotIn('CDT', headers(self.flight))
        self.assertIn('CHT 1', headers(self.flight))
        self.assertIsNone(next(rows(self.flight,self.download))[headers(self.flight).index('CHT 1')])

    def test_individual_alarm_thresholds(self):
        self.flight.series['C3'] = [401] * 5
        self.flight.series['T2'] = [1701] * 5
        self.flight.series['BAT'] = [23.9, 24, 30.5, 30.6, None]
        self.flight.series['CLD'] = [-201] * 5
        self.assertEqual(limits(self.flight, self.download, 0), 'TIT 2;CHT 3;BAT-L;CLD')
        self.assertNotIn('BAT', limits(self.flight, self.download, 1))
        self.assertNotIn('BAT', limits(self.flight, self.download, 2))
        self.assertIn('BAT-H', limits(self.flight, self.download, 3))
        self.assertNotIn('MAP', limits(self.flight, self.download, 0))

    def test_csv_xlsx_match_and_freeze(self):
        csv_path, xlsx_path = self.directory / 'out.csv', self.directory / 'out.xlsx'
        write_csv(csv_path, self.flight, self.download)
        write_xlsx(xlsx_path, [(self.flight,self.download)], [self.download], 0)
        with csv_path.open() as stream:
            table = list(csv.reader(stream))
        workbook = load_workbook(xlsx_path)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ['Summary', 'Flight 415'])
        for sheet in workbook:
            self.assertEqual(sheet.freeze_panes, 'C2')
            self.assertIsNone(sheet.auto_filter.ref)
            self.assertTrue(all(cell.comment is None for cell in sheet[1]))
        sheet = workbook['Flight 415']
        self.assertEqual(table[0], [c.value for c in sheet[1]])
        for csv_row, cells in zip(table[1:], sheet.iter_rows(min_row=2)):
            for index, (text, cell) in enumerate(zip(csv_row, cells)):
                value = cell.value
                if index == 2:
                    self.assertEqual(text, value.isoformat(sep=' '))
                elif value is None:
                    self.assertEqual(text, '')
                elif isinstance(value, (int, float)):
                    self.assertEqual(float(text), value)
                else:
                    self.assertEqual(text, value)
        self.assertEqual([r[0] for r in sheet.iter_rows(min_row=2, values_only=True)], [1,2,3,4,5])
        self.assertEqual([r[1] for r in sheet.iter_rows(min_row=2, values_only=True)], [0,2,4,6,8])
        self.assertEqual(sheet['C2'].number_format, 'yyyy-mm-dd hh:mm:ss')
        summary = workbook['Summary']
        self.assertEqual([cell.value for cell in summary[1]], ['Source', 'Property', 'Value'])
        for column, pixels in zip('ABC', [85,225,450]):
            self.assertAlmostEqual(summary.column_dimensions[column].width * 6, pixels, delta=1)
        self.assertTrue(all(cell.alignment.wrap_text for cell in summary['C']))
        download_cell = next(row[2] for row in summary if row[1].value == 'Download datetime')
        self.assertEqual(download_cell.number_format, 'yyyy-mm-dd hh:mm:ss')
        self.assertEqual(download_cell.value, self.download.metadata['Download datetime'])

    def test_widths_follow_data_when_columns_move(self):
        # Missing channels move Limits away from AE; navigation follows it.
        self.flight.series.pop('MAP')
        self.flight.series.pop('C1')
        self.flight.series['LAT'] = [None] * self.flight.samples
        target = self.directory / 'widths.xlsx'
        write_xlsx(target, [(self.flight,self.download)], [self.download], 0, summary=False)
        workbook = load_workbook(target)
        self.addCleanup(workbook.close)
        sheet = workbook['Flight 415']
        for cell in sheet[1]:
            narrow = {'FF', '% HP', 'OAT', 'CDT', 'IAT', 'Oil P', 'Oil T', 'Batt',
                      'Used', 'Mark', 'Diff', 'Cold'}
            expected = 53 if cell.value in narrow or cell.value.startswith('CHT ') else {
                'Sample':50,'Delta T':44,'DateTime':115,'Limits':125}.get(cell.value,55)
            self.assertAlmostEqual(sheet.column_dimensions[cell.column_letter].width * 6, expected, delta=1)
        self.assertIsNone(sheet.auto_filter.ref)
        self.assertTrue(all(cell.comment is None for cell in sheet[1]))

    def test_opening_window_geometry_in_combined_and_separate_workbooks(self):
        for visible_summary in (True, False):
            target = self.directory / f'window-{visible_summary}.xlsx'
            write_xlsx(target, [(self.flight,self.download)], [self.download], 0,
                       summary=visible_summary)
            workbook = load_workbook(target)
            self.addCleanup(workbook.close)
            view = workbook.views[0]
            self.assertEqual((view.xWindow, view.yWindow, view.windowWidth, view.windowHeight),
                             (4280, 2700, 32360, 18380))

    def test_metadata_is_literal_text(self):
        self.download.metadata['Aircraft ID'] = '=1+1'
        target = self.directory / 'literal.xlsx'
        write_xlsx(target, [(self.flight,self.download)], [self.download], 0)
        workbook = load_workbook(target)
        self.addCleanup(workbook.close)
        cell = next(row[2] for row in workbook['Summary'] if row[1].value == 'Aircraft ID')
        self.assertEqual(cell.value, '=1+1')
        self.assertEqual(cell.data_type, 's')

    def test_default_cutoff_and_explicit_short_flight_error(self):
        code, _, err = self.invoke('--csv', '--output-dir', self.directory, self.input)
        self.assertEqual(code, 0)
        self.assertIn('No flights', err)
        self.assertFalse(list(self.directory.glob('*.csv')))
        code, _, err = self.invoke('--csv', '--flights', '415', self.input)
        self.assertEqual(code, 201)
        self.assertIn('below minimum', err)

    def test_fractional_minute_cutoff(self):
        # Eight seconds is longer than 0.1 minutes but shorter than 0.2.
        code, _, _ = self.invoke('--csv','--minimum-duration','.1','--output-dir',self.directory,self.input)
        self.assertEqual(code, 0)
        self.assertEqual(len(list(self.directory.glob('*.csv'))), 1)

    def test_bad_file_does_not_block_good_file_or_status(self):
        bad = self.directory / 'bad.JPI'
        bad.write_bytes(synthetic_file().replace(b'U,TEST', b'U,FAIL'))
        code, _, err = self.invoke('--csv','--minimum-duration','0','--output-dir',self.directory,bad,self.input)
        self.assertEqual(code, errno.EILSEQ)
        self.assertTrue((self.directory / 'a_Flight_415.csv').exists())
        self.assertIn('bad.JPI', err)

    def test_explicit_selection_continues_after_corrupt_input(self):
        bad = self.directory / 'bad.JPI'
        bad.write_bytes(b'corrupt')
        code, _, err = self.invoke('--csv','--flights','414:415','--minimum-duration','0',
                                   '--output-dir',self.directory,bad,self.input)
        self.assertEqual(code, errno.EILSEQ)
        self.assertTrue((self.directory / 'a_Flight_415.csv').exists())
        self.assertIn('414', err)

    def test_no_partial_export_from_corrupt_file(self):
        bad = self.directory / 'bad.JPI'
        bad.write_bytes(synthetic_file()[:-9])
        code, _, _ = self.invoke('--csv','--minimum-duration','0','--output-dir',self.directory,bad)
        self.assertEqual(code, errno.EILSEQ)
        self.assertFalse(list(self.directory.glob('*.csv')))

    def test_duplicate_ids_preserved_in_combined_and_separate(self):
        other = self.directory / 'b.JPI'
        other.write_bytes(synthetic_file())
        combined = self.directory / 'combined.xlsx'
        code, _, err = self.invoke('--xls','--minimum-duration','0','--output',combined,self.input,other)
        self.assertEqual(code, 0)
        self.assertIn('preserving all copies', err)
        workbook = load_workbook(combined)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ['Summary','Flight 415','Flight 415 (2)'])
        code, _, _ = self.invoke('--xls-separate','--minimum-duration','0','--output-dir',self.directory,self.input,other)
        self.assertEqual(code, 0)
        for name in ['a_Flight_415.xlsx', 'b_Flight_415.xlsx']:
            book = load_workbook(self.directory / name)
            self.assertEqual(book.sheetnames, ['Summary', 'Flight 415'])
            self.assertEqual(book['Summary'].sheet_state, 'hidden')
            self.assertEqual(book.active.title, 'Flight 415')
            book.close()

    def test_existing_output_not_overwritten(self):
        target = self.directory / 'keep.csv'
        target.write_text('keep me')
        code, _, _ = self.invoke('--csv','--minimum-duration','0','--output',target,self.input)
        self.assertEqual(code, errno.EEXIST)
        self.assertEqual(target.read_text(), 'keep me')

    def test_multiple_outputs_reject_explicit_path(self):
        code, _, _ = self.invoke('--csv','--minimum-duration','0','--output',self.directory/'out.csv',self.input,self.input)
        self.assertEqual(code, errno.EINVAL)

    def test_selection_range_missing_id_and_engine(self):
        code, _, _ = self.invoke('--csv','--minimum-duration','0','--flights','414:415','--output-dir',self.directory,self.input)
        self.assertEqual(code, errno.EDOM)
        code, _, _ = self.invoke('--csv','--engine','1,3',self.input)
        self.assertEqual(code, errno.EDOM)
        code, _, _ = self.invoke('--csv','--minimum-duration','0','--engine','right','--flights','415:415','--output-dir',self.directory,self.input)
        self.assertEqual(code, 0)

    def test_info_and_list(self):
        code, out, _ = self.invoke('--info',self.input)
        self.assertEqual(code, 0)
        self.assertIn('BAT high alarm (V): 30.5', out)
        self.assertIn('Flight 415 cylinder count: 6', out)
        self.assertIn('Flight 415 IAT present: Yes', out)
        self.assertIn('Engine temperature units: F', out)
        code, out, _ = self.invoke('--list',self.input)
        self.assertEqual(code, 0)
        self.assertIn('00:00:08, No GPS', out)

    def test_info_selected_flight_keeps_file_metadata(self):
        code, out, err = self.invoke('--info','--flights','418',FIXTURES/'U260919.JPI')
        self.assertEqual(code,0,err)
        self.assertIn('Total flights: 4',out)
        self.assertIn('Flight 418 samples: 5507',out)
        for fid in (415,416,417):
            self.assertNotIn(f'Flight {fid} start:',out)

    def test_info_can_inspect_short_flight_and_select_engine(self):
        code, out, err = self.invoke('--info','--flights','415','--engine','right',self.input)
        self.assertEqual(code,0,err)
        self.assertIn('Flight 415 samples: 5',out)
        self.assertIn('Flights meeting minimum duration: 0',out)
        self.assertEqual(self.invoke('--info','--flights','999',self.input)[0],errno.EDOM)

    def test_info_with_list_or_validate(self):
        for action, message in [('--list','00:00:08, No GPS'),('--validate','a.JPI: Valid')]:
            with self.subTest(action=action):
                code, out, err = self.invoke(action,'--info','--flights','415',self.input)
                self.assertEqual(code,0,err)
                self.assertIn('Aircraft ID: TEST',out)
                self.assertIn(message,out)

    def test_info_with_each_export_in_either_order(self):
        for action in ('--csv','--xls','--xls-separate'):
            for first in (True,False):
                with self.subTest(action=action,info_first=first):
                    extension = '.csv' if action == '--csv' else '.xlsx'
                    target = self.directory / f'{action[2:]}-{first}{extension}'
                    flags = ['--info',action] if first else [action,'--info']
                    code, out, err = self.invoke(*flags,'--flights','415','--minimum-duration','0',
                                               '--output',target,self.input)
                    self.assertEqual(code,0,err)
                    self.assertIn('Flight 415 samples: 5',out)
                    self.assertTrue(target.exists())

    def test_info_with_export_keeps_duration_policy(self):
        code, out, err = self.invoke('--info','--csv','--flights','415',self.input)
        self.assertEqual(code,201)
        self.assertIn('Flight 415 samples: 5',out)
        self.assertIn('below minimum',err)

    def test_info_combination_preserves_bad_input_status(self):
        bad = self.directory / 'bad.JPI'
        bad.write_bytes(b'corrupt')
        code, out, err = self.invoke('--info','--csv','--flights','414:415',
                                   '--minimum-duration','0','--output-dir',self.directory,bad,self.input)
        self.assertEqual(code,errno.EILSEQ)
        self.assertIn('Flight 415 samples: 5',out)
        self.assertTrue((self.directory / 'a_Flight_415.csv').exists())

    def test_info_with_help_or_version(self):
        for action, message in [('--help','usage:'),('--version','JPI2Excel')]:
            with self.subTest(action=action):
                code, out, err = self.invoke('--info',action,self.input)
                self.assertEqual(code,0,err)
                self.assertIn(message,out)
                self.assertIn('Aircraft ID: TEST',out)
                self.assertEqual(self.invoke('--info',action)[0],0)

    def test_info_does_not_allow_output_paths_without_export(self):
        self.assertEqual(self.invoke('--info','--output',self.directory/'out.csv',self.input)[0],errno.EINVAL)

    def test_argument_errors_help_and_version(self):
        for args, expected in [([],errno.EOPNOTSUPP),(['--bad'],errno.EINVAL),
                               (['--csv','--xls'],errno.EINVAL),
                               (['--info','--csv','--xls'],errno.EINVAL),
                               (['--csv','--minimum-duration','nan',str(self.input)],errno.EINVAL),
                               (['--help'],0),(['--version'],0)]:
            with self.subTest(args=args):
                self.assertEqual(self.invoke(*args)[0], expected)
        help_text = self.invoke('--help')[1]
        self.assertIn('JPI2Excel.xlsx',help_text)
        self.assertIn('<source>_Flight_<number>',help_text)

    def test_destination_errors(self):
        for target, expected in [(self.directory,errno.EISDIR),
                                 (self.directory/'absent'/'out.csv',errno.ENOENT)]:
            self.assertEqual(self.invoke('--csv','--minimum-duration','0','--output',target,self.input)[0],expected)

    def test_cancel_returns_ecanceled(self):
        with patch('jpi2excel.reader.read_jpi', side_effect=OSError(errno.ECANCELED,'interrupted')):
            self.assertEqual(self.invoke('--validate',self.input)[0],errno.ECANCELED)

    def test_real_fixture_block_slack_accepted(self):
        # Validated physical block slack is not corrupt flight data.
        for name in ['U260731.JPI','U260828.JPI','U260919.JPI']:
            code, out, err = self.invoke('--validate',FIXTURES/name)
            self.assertEqual(code,0,err)
            self.assertIn(f'{name}: Valid',out)

    def test_original_fixture_exports_exclude_slack(self):
        expected = {
            'U260731': {372:4872, 373:3547},
            'U260828': {393:1430, 394:1470, 395:885, 396:679,
                       397:1348, 398:1299, 399:1683, 400:1209},
            'U260919': {415:724, 416:1008, 417:33, 418:5507},
        }
        inputs = [FIXTURES / (stem + '.JPI') for stem in expected]
        code, _, err = self.invoke('--csv','--minimum-duration','0',
                                   '--output-dir',self.directory,*inputs)
        self.assertEqual(code,0,err)
        self.assertEqual(len(list(self.directory.glob('*.csv'))),14)
        for stem, flights in expected.items():
            for fid, samples in flights.items():
                with (self.directory / f'{stem}_Flight_{fid}.csv').open() as stream:
                    table = list(csv.DictReader(stream))
                self.assertEqual(len(table),samples)
                self.assertEqual(int(table[-1]['Sample']),samples)
        target = self.directory / 'original.xlsx'
        code, _, err = self.invoke('--xls','--minimum-duration','0',
                                   '--output',target,inputs[-1])
        self.assertEqual(code,0,err)
        workbook = load_workbook(target,read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames,['Summary','Flight 415','Flight 416','Flight 417','Flight 418'])
        for fid, samples in expected['U260919'].items():
            self.assertEqual(workbook[f'Flight {fid}'].max_row,samples + 1)
        properties = {row[1]: row[2] for row in workbook['Summary'].iter_rows(min_row=2,values_only=True)}
        self.assertEqual(properties['Binary block count ($L)'],369)
        self.assertEqual(properties['Block slack length (bytes)'],224)
        self.assertEqual(properties['Block slack nonzero'],True)
        code, out, err = self.invoke('--info',inputs[-1])
        self.assertEqual(code,0,err)
        self.assertIn('Block slack length (bytes): 224',out)
